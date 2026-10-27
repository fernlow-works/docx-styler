"""Checks one exported .docx: real Word styles and every Markdown feature. Prints JSON, exits 1 on failure."""
import json, sys, zipfile
import docx

path = sys.argv[1]
d = docx.Document(path)
styles = [p.style.name for p in d.paragraphs]
text = "\n".join(p.text for p in d.paragraphs)
with zipfile.ZipFile(path) as z:
    names = z.namelist()
    doc_xml = z.read("word/document.xml").decode()
    foot = z.read("word/footnotes.xml").decode() if "word/footnotes.xml" in names else ""
    media = [n for n in names if n.startswith("word/media/") and not n.endswith("/")]

checks = {
    "title_style": "Title" in styles,
    "heading_1_2_3_real_styles": all(h in styles for h in ("Heading 1", "Heading 2", "Heading 3")),
    "heading_style_builtin": all(p.style.builtin for p in d.paragraphs if p.style.name.startswith("Heading")),
    "list_paragraphs": styles.count("List Paragraph") >= 7,
    "numbered_list": "<w:numPr>" in doc_xml,
    "tasks": "\u2610" in text and "\u2612" in text,
    "callout_and_quote_style": styles.count("Quote") >= 2 and "Remember this" in text,
    "table": len(d.tables) == 1 and d.tables[0].cell(1, 1).text == "b1",
    "code_block_style": "DocxStylerCode" in [p.style.style_id for p in d.paragraphs],
    "image_embedded": len(media) == 1 and len(d.inline_shapes) == 1,
    "footnote": "The footnote text." in foot and "<w:footnoteReference" in doc_xml,
    "hyperlink": "https://example.com" in json.dumps([r.target_ref for r in d.part.rels.values()]),
    "frontmatter_removed": "tags:" not in text,
    "comment_removed": "hidden comment" not in text,
    "wikilink_alias": "alias" in text and "[[" not in text,
}
print(json.dumps({"file": path, "paragraphs": len(d.paragraphs), "checks": checks}))
sys.exit(0 if all(checks.values()) else 1)
